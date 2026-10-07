"""Approximate text-overflow QA for a .pptx when no renderer (LibreOffice/
PowerPoint) is available in the environment. Measures each text box's actual
runs with PIL against the real Calibri font metrics and estimates wrapped
line count vs. the box's height. Not a substitute for a real render -- flags
anything that looks close or over so it can be fixed conservatively.

Usage: python scripts/qa_overflow_check.py docs/progress_update.pptx
"""
import sys

from pptx import Presentation
from pptx.util import Emu
from PIL import ImageFont

EMU_PER_PT = 12700
PX_PER_PT = 96 / 72  # assume 96 DPI, matches PowerPoint's default on-screen metric
MARGIN_L = MARGIN_R = 91440   # python-pptx default internal inset, EMU (0.1")
MARGIN_T = MARGIN_B = 45720   # 0.05"

FONT_PATHS = {
    (False, False): "C:/Windows/Fonts/calibri.ttf",
    (True, False): "C:/Windows/Fonts/calibrib.ttf",
    (False, True): "C:/Windows/Fonts/calibrii.ttf",
    (True, True): "C:/Windows/Fonts/calibriz.ttf",
}
_font_cache = {}


def get_font(size_pt, bold, italic):
    key = (round(size_pt * 4), bold, italic)
    if key not in _font_cache:
        path = FONT_PATHS.get((bold, italic), FONT_PATHS[(False, False)])
        px_size = max(1, round(size_pt * PX_PER_PT))
        _font_cache[key] = ImageFont.truetype(path, px_size)
    return _font_cache[key]


def wrap_lines(text, font, max_width_px):
    if not text:
        return 1
    words = text.split(" ")
    lines, cur = 1, ""
    for w in words:
        trial = (cur + " " + w).strip()
        width = font.getbbox(trial)[2]
        if width > max_width_px and cur:
            lines += 1
            cur = w
        else:
            cur = trial
    return lines


def check_shape(shape, slide_idx):
    if not shape.has_text_frame:
        return []
    tf = shape.text_frame
    usable_w_px = max(1, (shape.width - MARGIN_L - MARGIN_R)) / EMU_PER_PT * PX_PER_PT
    issues = []
    total_h_pt = 0.0
    for p in tf.paragraphs:
        line_spacing = p.line_spacing if isinstance(p.line_spacing, (int, float)) else 1.0
        space_after_pt = (p.space_after.pt if p.space_after else 0)
        text = "".join(r.text for r in p.runs)
        if not text.strip():
            continue
        size_pt = None
        bold = italic = False
        for r in p.runs:
            if r.font.size:
                size_pt = r.font.size.pt
            bold = bold or bool(r.font.bold)
            italic = italic or bool(r.font.italic)
        size_pt = size_pt or 14
        font = get_font(size_pt, bold, italic)
        n_lines = wrap_lines(text, font, usable_w_px)
        line_h_pt = size_pt * line_spacing * 1.2
        total_h_pt += n_lines * line_h_pt + space_after_pt
    usable_h_pt = max(1, (shape.height - MARGIN_T - MARGIN_B)) / EMU_PER_PT
    if total_h_pt > usable_h_pt:
        ratio = total_h_pt / usable_h_pt
        issues.append(
            f"slide {slide_idx}: shape at ({Emu(shape.left).inches:.2f}in,"
            f"{Emu(shape.top).inches:.2f}in) size {Emu(shape.width).inches:.2f}x"
            f"{Emu(shape.height).inches:.2f}in -- est. content {total_h_pt:.0f}pt "
            f"vs box {usable_h_pt:.0f}pt ({ratio:.2f}x) -- "
            f"{'LIKELY OVERFLOW' if ratio > 1.08 else 'tight, check manually'}"
        )
    return issues


def main(path):
    prs = Presentation(path)
    all_issues = []
    for i, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            all_issues.extend(check_shape(shape, i))
    if not all_issues:
        print("No overflow risk flagged by this approximate check.")
    else:
        print(f"{len(all_issues)} shape(s) flagged:")
        for issue in all_issues:
            print(" -", issue)


if __name__ == "__main__":
    main(sys.argv[1])
